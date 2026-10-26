"""Markdown из docs/ → Word для чтения владельцем.
python3 tools/docs/md2docx.py docs/PLAN.md docs/PLAN.docx "4=1.6,7.2,2.0,6.2" "2=3.0,14.0"
Ширины колонок в см по числу колонок таблицы; без указания — поровну. Нужны pandoc и python-docx."""
import sys, subprocess, tempfile, re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
W = {}
for a in sys.argv[3:]:
    n, ws = a.split('='); W[int(n)] = [float(x) for x in ws.split(',')]

# список сразу после строки текста без пустой строки pandoc склеивает в абзац
lines = open(src, encoding='utf-8').read().split('\n'); out = []
for l in lines:
    if re.match(r'^(- |\d+\. )', l) and out and out[-1].strip() and not re.match(r'^(- |\d+\. |  )', out[-1]):
        out.append('')
    out.append(l)
tmp = tempfile.mkdtemp()
open(tmp + '/in.md', 'w', encoding='utf-8').write('\n'.join(out))
subprocess.run(['pandoc', tmp + '/in.md', '-f', 'markdown', '-t', 'docx', '-o', tmp + '/raw.docx'], check=True)

d = docx.Document(tmp + '/raw.docx')
for sec in d.sections:
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
by = {s.name.lower(): s for s in d.styles}
for st in d.styles:
    try:
        rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), 'Calibri')
        for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
    except Exception: pass
for n in ('normal', 'body text', 'first paragraph', 'compact'):
    if n in by: by[n].font.size = Pt(10.5)
for n, sz in (('heading 1', 20), ('heading 2', 15), ('heading 3', 12.5), ('title', 22)):
    if n in by:
        st = by[n]; st.font.size = Pt(sz); st.font.bold = True; st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
for t in d.tables:
    n = len(t.columns); w = W.get(n) or [17.0 / n] * n
    tblPr = t._tbl.tblPr
    for tag in ('w:tblW', 'w:tblLayout', 'w:tblBorders'):
        e = tblPr.find(qn(tag))
        if e is not None: tblPr.remove(e)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(w) * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), 'B8BFCC'); b.append(x)
    tblPr.append(b)
    for gc, wi in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), w): gc.set(qn('w:w'), str(int(wi * 567)))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            c.width = Cm(w[ci])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if ri == 0: r.font.bold = True
            if ri == 0:
                tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'E8ECF3'); tcPr.append(sh)
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
d.save(dst)
print(dst, len(d.tables), 'таблиц')
